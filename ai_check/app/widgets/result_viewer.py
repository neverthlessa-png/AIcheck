"""结果查看器组件"""

from pathlib import Path
from typing import Any

from loguru import logger
from PyQt6.QtCore import Qt
from PyQt6.QtGui import QColor
from PyQt6.QtWidgets import (
    QAbstractItemView,
    QFileDialog,
    QGroupBox,
    QHeaderView,
    QLabel,
    QSplitter,
    QTableWidget,
    QTableWidgetItem,
    QTextEdit,
    QVBoxLayout,
    QWidget,
)


class ResultViewer(QWidget):
    """结果查看器"""

    def __init__(self) -> None:
        super().__init__()
        self._results: list[dict[str, Any]] = []
        self._init_ui()

    def _init_ui(self) -> None:
        """初始化UI"""
        layout = QVBoxLayout(self)

        # 统计信息
        self._stats_group = QGroupBox("统计信息")
        stats_layout = QVBoxLayout(self._stats_group)

        self._stats_label = QLabel("检测结果: 共 0 张图片, 0 张异常")
        stats_layout.addWidget(self._stats_label)

        layout.addWidget(self._stats_group)

        # 分割器
        splitter = QSplitter(Qt.Orientation.Vertical)

        # 结果表格
        self._result_table = QTableWidget()
        self._result_table.setColumnCount(6)
        self._result_table.setHorizontalHeaderLabels(
            ["图片", "检测器", "类型", "异常", "置信度", "描述"]
        )
        self._result_table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self._result_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)

        header = self._result_table.horizontalHeader()
        header.setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
        header.setSectionResizeMode(5, QHeaderView.ResizeMode.Stretch)

        splitter.addWidget(self._result_table)

        # 详情面板
        self._detail_panel = QTextEdit()
        self._detail_panel.setReadOnly(True)
        self._detail_panel.setMaximumHeight(150)
        splitter.addWidget(self._detail_panel)

        layout.addWidget(splitter)

        # 连接信号
        self._result_table.itemSelectionChanged.connect(self._on_selection_changed)

    def add_result(self, result: dict[str, Any]) -> None:
        """添加检测结果"""
        self._results.append(result)
        self._update_table()

    def set_results(self, results: list[dict[str, Any]]) -> None:
        """设置检测结果"""
        self._results = results
        self._update_table()

    def _update_table(self) -> None:
        """更新表格"""
        self._result_table.setRowCount(0)

        for result in self._results:
            image_path = Path(result.get("image_path", ""))
            for det_result in result.get("results", []):
                row = self._result_table.rowCount()
                self._result_table.insertRow(row)

                # 图片名称
                self._result_table.setItem(row, 0, QTableWidgetItem(image_path.name))

                # 检测器名称
                self._result_table.setItem(
                    row, 1, QTableWidgetItem(det_result.get("detector_name", ""))
                )

                # 检测类型
                self._result_table.setItem(
                    row, 2, QTableWidgetItem(det_result.get("detection_type", ""))
                )

                # 是否异常
                is_anomaly = det_result.get("is_anomaly", False)
                anomaly_item = QTableWidgetItem("是" if is_anomaly else "否")
                if is_anomaly:
                    anomaly_item.setBackground(QColor(255, 200, 200))
                else:
                    anomaly_item.setBackground(QColor(200, 255, 200))
                self._result_table.setItem(row, 3, anomaly_item)

                # 置信度
                confidence = det_result.get("confidence", 0)
                self._result_table.setItem(row, 4, QTableWidgetItem(f"{confidence:.2%}"))

                # 描述
                self._result_table.setItem(
                    row, 5, QTableWidgetItem(det_result.get("description", ""))
                )

        # 更新统计
        total = len(self._results)
        anomalies = sum(1 for r in self._results if r.get("has_anomaly", False))
        self._stats_label.setText(f"检测结果: 共 {total} 张图片, {anomalies} 张异常")

    def _on_selection_changed(self) -> None:
        """选择变化"""
        selected_rows = self._result_table.selectedItems()
        if not selected_rows:
            return

        row = selected_rows[0].row()
        if row < len(self._results):
            result = self._results[row]
            self._show_detail(result)

    def _show_detail(self, result: dict[str, Any]) -> None:
        """显示详情"""
        import json

        detail_text = f"图片: {result.get('image_path', '')}\n"
        detail_text += f"异常: {result.get('has_anomaly', False)}\n"
        detail_text += f"最高置信度: {result.get('max_confidence', 0):.2%}\n\n"

        for det_result in result.get("results", []):
            detail_text += f"--- {det_result.get('detector_name', '')} ---\n"
            detail_text += f"类型: {det_result.get('detection_type', '')}\n"
            detail_text += f"异常级别: {det_result.get('anomaly_level', '')}\n"
            detail_text += f"置信度: {det_result.get('confidence', 0):.2%}\n"
            detail_text += f"描述: {det_result.get('description', '')}\n"

            metadata = det_result.get("metadata", {})
            if metadata:
                detail_text += f"元数据: {json.dumps(metadata, indent=2)}\n"

            detail_text += "\n"

        self._detail_panel.setText(detail_text)

    def export_report(self, path: Path) -> None:
        """导出报告"""
        try:
            suffix = path.suffix.lower()

            if suffix == ".xlsx":
                self._export_excel(path)
            elif suffix == ".csv":
                self._export_csv(path)
            else:
                logger.error(f"不支持的格式: {suffix}")
                return

            logger.info(f"报告已导出: {path}")

        except Exception as e:
            logger.error(f"导出报告失败: {e}")

    def _export_excel(self, path: Path) -> None:
        """导出Excel"""
        import openpyxl
        from openpyxl.styles import Font, PatternFill

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "检测结果"

        # 表头
        headers = ["图片", "检测器", "类型", "异常", "置信度", "描述"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")

        # 数据
        row_num = 2
        for result in self._results:
            image_path = Path(result.get("image_path", ""))
            for det_result in result.get("results", []):
                ws.cell(row=row_num, column=1, value=image_path.name)
                ws.cell(row=row_num, column=2, value=det_result.get("detector_name", ""))
                ws.cell(row=row_num, column=3, value=det_result.get("detection_type", ""))
                ws.cell(row=row_num, column=4, value="是" if det_result.get("is_anomaly") else "否")
                ws.cell(row=row_num, column=5, value=det_result.get("confidence", 0))
                ws.cell(row=row_num, column=6, value=det_result.get("description", ""))
                row_num += 1

        wb.save(path)

    def _export_csv(self, path: Path) -> None:
        """导出CSV"""
        import csv

        with open(path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)

            # 表头
            writer.writerow(["图片", "检测器", "类型", "异常", "置信度", "描述"])

            # 数据
            for result in self._results:
                image_path = Path(result.get("image_path", ""))
                for det_result in result.get("results", []):
                    writer.writerow(
                        [
                            image_path.name,
                            det_result.get("detector_name", ""),
                            det_result.get("detection_type", ""),
                            "是" if det_result.get("is_anomaly") else "否",
                            det_result.get("confidence", 0),
                            det_result.get("description", ""),
                        ]
                    )

    def clear(self) -> None:
        """清空结果"""
        self._results.clear()
        self._result_table.setRowCount(0)
        self._stats_label.setText("检测结果: 共 0 张图片, 0 张异常")
        self._detail_panel.clear()
